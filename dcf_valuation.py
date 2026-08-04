import argparse
import sys
from datetime import datetime

import numpy as np
import pandas as pd
import yfinance as yf
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, Reference


# --------------------------------------------------------------------------- #
# Core DCF engine
# --------------------------------------------------------------------------- #

class DCFModel:
    def __init__(self, ticker, projection_years=5, fcf_growth_rate=None,
                 terminal_growth_rate=0.025, risk_free_rate=0.043,
                 market_risk_premium=0.055, tax_rate=None):
        self.ticker_symbol = ticker.upper()
        self.projection_years = projection_years
        self.fcf_growth_rate = fcf_growth_rate      # None => auto-estimated
        self.terminal_growth_rate = terminal_growth_rate
        self.risk_free_rate = risk_free_rate
        self.market_risk_premium = market_risk_premium
        self.tax_rate = tax_rate                    # None => auto-estimated
        self.warnings = []
        self.data = {}
        self.results = {}

    # ------------------------------------------------------------------ #
    # 1. Data retrieval
    # ------------------------------------------------------------------ #
    def fetch_data(self):
        print(f"Fetching live data for {self.ticker_symbol} ...")
        tk = yf.Ticker(self.ticker_symbol)
        info = tk.info
        if not info or info.get("regularMarketPrice") is None and info.get("currentPrice") is None:
            raise ValueError(f"No market data returned for '{self.ticker_symbol}'. "
                              f"Check the ticker symbol.")

        cashflow = tk.cashflow            # annual, most recent first
        balance_sheet = tk.balance_sheet
        financials = tk.financials

        self.data["info"] = info
        self.data["cashflow"] = cashflow
        self.data["balance_sheet"] = balance_sheet
        self.data["financials"] = financials

        # --- Free cash flow history (Operating CF - CapEx) ---
        fcf_history = self._compute_historical_fcf(cashflow)
        if fcf_history.empty:
            raise ValueError("Could not derive historical free cash flow from statements.")
        self.data["fcf_history"] = fcf_history

        # --- Shares outstanding / price / market cap ---
        self.data["shares_out"] = info.get("sharesOutstanding")
        self.data["price"] = info.get("currentPrice") or info.get("regularMarketPrice")
        self.data["market_cap"] = info.get("marketCap")
        self.data["beta"] = info.get("beta") or 1.0
        if info.get("beta") is None:
            self.warnings.append("Beta not available; defaulted to 1.0.")

        # --- Balance sheet items for WACC / net debt ---
        self.data["total_debt"] = self._latest(balance_sheet, [
            "Total Debt", "Long Term Debt", "Short Long Term Debt"])
        if self.data["total_debt"] is None:
            self.data["total_debt"] = info.get("totalDebt") or 0
            self.warnings.append("Total debt pulled from summary info (statement lookup failed).")

        self.data["cash"] = self._latest(balance_sheet, [
            "Cash And Cash Equivalents", "Cash Cash Equivalents And Short Term Investments"])
        if self.data["cash"] is None:
            self.data["cash"] = info.get("totalCash") or 0
            self.warnings.append("Cash pulled from summary info (statement lookup failed).")

        # --- Tax rate ---
        if self.tax_rate is None:
            self.tax_rate = self._estimate_tax_rate(financials)

        # --- Cost of debt (interest expense / total debt, fallback 4.5%) ---
        self.data["cost_of_debt"] = self._estimate_cost_of_debt(financials)

        print("Data retrieval complete.\n")
        return self.data

    @staticmethod
    def _latest(df, row_candidates):
        if df is None or df.empty:
            return None
        for name in row_candidates:
            if name in df.index:
                val = df.loc[name].dropna()
                if not val.empty:
                    return float(val.iloc[0])
        return None

    def _compute_historical_fcf(self, cashflow):
        if cashflow is None or cashflow.empty:
            return pd.Series(dtype=float)

        op_cf_row = None
        for name in ["Operating Cash Flow", "Total Cash From Operating Activities"]:
            if name in cashflow.index:
                op_cf_row = cashflow.loc[name]
                break

        capex_row = None
        for name in ["Capital Expenditure", "Capital Expenditures"]:
            if name in cashflow.index:
                capex_row = cashflow.loc[name]
                break

        if op_cf_row is None or capex_row is None:
            return pd.Series(dtype=float)

        fcf = (op_cf_row + capex_row).dropna()  # capex is stored negative
        fcf = fcf.sort_index()  # oldest -> newest
        return fcf

    def _estimate_tax_rate(self, financials):
        try:
            pretax = None
            tax_exp = None
            for name in ["Pretax Income", "Income Before Tax"]:
                if name in financials.index:
                    pretax = financials.loc[name].dropna().iloc[0]
                    break
            for name in ["Tax Provision", "Income Tax Expense"]:
                if name in financials.index:
                    tax_exp = financials.loc[name].dropna().iloc[0]
                    break
            if pretax and tax_exp and pretax != 0:
                rate = float(tax_exp / pretax)
                if 0 < rate < 0.5:
                    return rate
        except Exception:
            pass
        self.warnings.append("Effective tax rate not derivable from statements; defaulted to 21%.")
        return 0.21

    def _estimate_cost_of_debt(self, financials):
        try:
            interest_exp = None
            for name in ["Interest Expense"]:
                if name in financials.index:
                    interest_exp = abs(financials.loc[name].dropna().iloc[0])
                    break
            debt = self.data.get("total_debt") or 0
            if interest_exp and debt:
                return float(interest_exp / debt)
        except Exception:
            pass
        self.warnings.append("Cost of debt not derivable from statements; defaulted to 4.5%.")
        return 0.045

    # ------------------------------------------------------------------ #
    # 2. WACC
    # ------------------------------------------------------------------ #
    def calculate_wacc(self):
        beta = self.data["beta"]
        cost_of_equity = self.risk_free_rate + beta * self.market_risk_premium

        market_cap = self.data["market_cap"] or 0
        debt = self.data["total_debt"] or 0
        total_capital = market_cap + debt

        if total_capital == 0:
            wacc = cost_of_equity
            self.warnings.append("Could not compute capital structure weights; WACC set to cost of equity.")
        else:
            weight_equity = market_cap / total_capital
            weight_debt = debt / total_capital
            after_tax_cod = self.data["cost_of_debt"] * (1 - self.tax_rate)
            wacc = weight_equity * cost_of_equity + weight_debt * after_tax_cod

        # Sanity guard rails
        wacc = max(min(wacc, 0.20), 0.03)

        self.results["cost_of_equity"] = cost_of_equity
        self.results["wacc"] = wacc
        return wacc

    # ------------------------------------------------------------------ #
    # 3. FCF projection
    # ------------------------------------------------------------------ #
    def project_fcf(self):
        fcf_hist = self.data["fcf_history"]
        base_fcf = float(fcf_hist.iloc[-1])

        if self.fcf_growth_rate is None:
            growth = self._estimate_growth_rate(fcf_hist)
            self.fcf_growth_rate = growth
        else:
            growth = self.fcf_growth_rate

        projections = []
        fcf = base_fcf
        for year in range(1, self.projection_years + 1):
            fcf = fcf * (1 + growth)
            projections.append(fcf)

        self.results["base_fcf"] = base_fcf
        self.results["projected_fcf"] = projections
        self.results["growth_rate_used"] = growth
        return projections

    def _estimate_growth_rate(self, fcf_hist):
        clean = fcf_hist[fcf_hist > 0]
        if len(clean) >= 2:
            n = len(clean) - 1
            try:
                cagr = (clean.iloc[-1] / clean.iloc[0]) ** (1 / n) - 1
                # Cap to a believable range
                cagr = max(min(cagr, 0.25), -0.10)
                return cagr
            except Exception:
                pass
        self.warnings.append("Historical FCF growth rate not reliable; defaulted to 5%.")
        return 0.05

    # ------------------------------------------------------------------ #
    # 4. Discounting, terminal value, equity bridge
    # ------------------------------------------------------------------ #
    def discount_cash_flows(self):
        wacc = self.results["wacc"]
        projections = self.results["projected_fcf"]

        discounted = [cf / ((1 + wacc) ** (i + 1)) for i, cf in enumerate(projections)]
        pv_fcf_sum = sum(discounted)

        terminal_value = (projections[-1] * (1 + self.terminal_growth_rate) /
                           (wacc - self.terminal_growth_rate))
        pv_terminal_value = terminal_value / ((1 + wacc) ** self.projection_years)

        enterprise_value = pv_fcf_sum + pv_terminal_value

        cash = self.data["cash"] or 0
        debt = self.data["total_debt"] or 0
        equity_value = enterprise_value + cash - debt

        shares_out = self.data["shares_out"]
        intrinsic_value_per_share = equity_value / shares_out if shares_out else None

        self.results["discounted_fcf"] = discounted
        self.results["pv_fcf_sum"] = pv_fcf_sum
        self.results["terminal_value"] = terminal_value
        self.results["pv_terminal_value"] = pv_terminal_value
        self.results["enterprise_value"] = enterprise_value
        self.results["equity_value"] = equity_value
        self.results["intrinsic_value_per_share"] = intrinsic_value_per_share
        return self.results

    # ------------------------------------------------------------------ #
    # 5. Orchestration
    # ------------------------------------------------------------------ #
    def run(self):
        self.fetch_data()
        self.calculate_wacc()
        self.project_fcf()
        self.discount_cash_flows()
        self._summarize()
        return self.results

    def _summarize(self):
        price = self.data["price"]
        ivps = self.results["intrinsic_value_per_share"]
        if price and ivps:
            upside = (ivps / price) - 1
            self.results["upside_pct"] = upside
            verdict = "UNDERVALUED" if upside > 0 else "OVERVALUED"
            self.results["verdict"] = verdict

        print(f"--- {self.ticker_symbol} DCF Summary ---")
        print(f"Current price:            {price:,.2f}" if price else "Current price:            n/a")
        print(f"WACC:                     {self.results['wacc']*100:.2f}%")
        print(f"FCF growth rate used:     {self.results['growth_rate_used']*100:.2f}%")
        print(f"Terminal growth rate:     {self.terminal_growth_rate*100:.2f}%")
        print(f"Enterprise value:         {self.results['enterprise_value']:,.0f}")
        print(f"Equity value:             {self.results['equity_value']:,.0f}")
        if ivps:
            print(f"Intrinsic value / share:  {ivps:,.2f}")
            print(f"Implied upside/downside:  {self.results['upside_pct']*100:.1f}%  ({self.results.get('verdict')})")
        if self.warnings:
            print("\nWarnings / fallbacks used:")
            for w in self.warnings:
                print(f"  - {w}")
        print()


# --------------------------------------------------------------------------- #
# Excel export
# --------------------------------------------------------------------------- #

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
SUBHEADER_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=12)
BOLD = Font(bold=True)
TITLE_FONT = Font(bold=True, size=16, color="1F4E78")
THIN = Side(style="thin", color="B7B7B7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _style_header_row(ws, row, n_cols):
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        cell.border = BORDER


def _autofit(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def export_to_excel(model: DCFModel, filename: str):
    wb = Workbook()

    # ---------------- Summary sheet ----------------
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"{model.ticker_symbol} — DCF Valuation Summary"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws["A2"].font = Font(italic=True, size=9, color="808080")

    r = model.results
    d = model.data
    rows = [
        ("Current share price", d.get("price")),
        ("Shares outstanding", d.get("shares_out")),
        ("Market capitalization", d.get("market_cap")),
        ("", ""),
        ("WACC", r.get("wacc")),
        ("Cost of equity", r.get("cost_of_equity")),
        ("Tax rate (effective)", model.tax_rate),
        ("FCF growth rate (projection)", r.get("growth_rate_used")),
        ("Terminal growth rate", model.terminal_growth_rate),
        ("", ""),
        ("Base year FCF", r.get("base_fcf")),
        ("PV of projected FCF", r.get("pv_fcf_sum")),
        ("Terminal value (undiscounted)", r.get("terminal_value")),
        ("PV of terminal value", r.get("pv_terminal_value")),
        ("Enterprise value", r.get("enterprise_value")),
        ("(+) Cash & equivalents", d.get("cash")),
        ("(-) Total debt", d.get("total_debt")),
        ("Equity value", r.get("equity_value")),
        ("", ""),
        ("Intrinsic value per share", r.get("intrinsic_value_per_share")),
        ("Implied upside / (downside)", r.get("upside_pct")),
        ("Verdict", r.get("verdict")),
    ]

    start_row = 4
    ws.cell(row=start_row, column=1, value="Metric").font = BOLD
    ws.cell(row=start_row, column=2, value="Value").font = BOLD
    _style_header_row(ws, start_row, 2)

    pct_labels = {"WACC", "Cost of equity", "Tax rate (effective)", "FCF growth rate (projection)",
                  "Terminal growth rate", "Implied upside / (downside)"}
    money_labels = {"Current share price", "Market capitalization", "Base year FCF",
                    "PV of projected FCF", "Terminal value (undiscounted)", "PV of terminal value",
                    "Enterprise value", "(+) Cash & equivalents", "(-) Total debt", "Equity value",
                    "Intrinsic value per share"}

    for i, (label, value) in enumerate(rows, start=start_row + 1):
        c_label = ws.cell(row=i, column=1, value=label)
        c_val = ws.cell(row=i, column=2, value=value)
        c_label.border = BORDER
        c_val.border = BORDER
        if label in {"Enterprise value", "Equity value", "Intrinsic value per share", "Verdict"}:
            c_label.font = BOLD
            c_val.font = BOLD
        if label in pct_labels and isinstance(value, (int, float)):
            c_val.number_format = "0.00%"
        elif label in money_labels and isinstance(value, (int, float)):
            c_val.number_format = "#,##0.00"
        elif label == "Shares outstanding" and isinstance(value, (int, float)):
            c_val.number_format = "#,##0"

    _autofit(ws, [32, 20])

    if model.warnings:
        wr = start_row + len(rows) + 3
        ws.cell(row=wr, column=1, value="Notes / fallback assumptions used:").font = BOLD
        for i, w in enumerate(model.warnings, start=1):
            ws.cell(row=wr + i, column=1, value=f"- {w}")

    # ---------------- Projections sheet ----------------
    ws2 = wb.create_sheet("FCF Projection")
    ws2["A1"] = f"{model.ticker_symbol} — Free Cash Flow Projection"
    ws2["A1"].font = TITLE_FONT

    headers = ["Year", "Projected FCF", "Discount Factor", "PV of FCF"]
    hr = 3
    for c, h in enumerate(headers, start=1):
        ws2.cell(row=hr, column=c, value=h)
    _style_header_row(ws2, hr, len(headers))

    wacc = r["wacc"]
    for i, (fcf, pv) in enumerate(zip(r["projected_fcf"], r["discounted_fcf"]), start=1):
        row = hr + i
        disc_factor = 1 / ((1 + wacc) ** i)
        ws2.cell(row=row, column=1, value=f"Year {i}").border = BORDER
        c = ws2.cell(row=row, column=2, value=fcf); c.number_format = "#,##0.00"; c.border = BORDER
        c = ws2.cell(row=row, column=3, value=disc_factor); c.number_format = "0.0000"; c.border = BORDER
        c = ws2.cell(row=row, column=4, value=pv); c.number_format = "#,##0.00"; c.border = BORDER

    total_row = hr + len(r["projected_fcf"]) + 1
    ws2.cell(row=total_row, column=1, value="Sum of PV(FCF)").font = BOLD
    c = ws2.cell(row=total_row, column=4, value=r["pv_fcf_sum"])
    c.number_format = "#,##0.00"
    c.font = BOLD

    _autofit(ws2, [16, 20, 18, 20])

    # Simple chart of projected FCF
    chart = LineChart()
    chart.title = "Projected Free Cash Flow"
    chart.y_axis.title = "FCF"
    chart.x_axis.title = "Year"
    data_ref = Reference(ws2, min_col=2, min_row=hr, max_row=hr + len(r["projected_fcf"]))
    cats_ref = Reference(ws2, min_col=1, min_row=hr + 1, max_row=hr + len(r["projected_fcf"]))
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    ws2.add_chart(chart, "F3")

    # ---------------- Historical financials sheet ----------------
    ws3 = wb.create_sheet("Historical Data")
    ws3["A1"] = f"{model.ticker_symbol} — Historical Free Cash Flow"
    ws3["A1"].font = TITLE_FONT
    ws3.cell(row=3, column=1, value="Fiscal Period End")
    ws3.cell(row=3, column=2, value="Free Cash Flow")
    _style_header_row(ws3, 3, 2)
    fcf_hist = d["fcf_history"]
    for i, (idx, val) in enumerate(fcf_hist.items(), start=1):
        date_label = idx.strftime("%Y-%m-%d") if hasattr(idx, "strftime") else str(idx)
        ws3.cell(row=3 + i, column=1, value=date_label).border = BORDER
        c = ws3.cell(row=3 + i, column=2, value=float(val))
        c.number_format = "#,##0.00"
        c.border = BORDER
    _autofit(ws3, [22, 20])

    # ---------------- Assumptions sheet ----------------
    ws4 = wb.create_sheet("Assumptions")
    ws4["A1"] = "Model Assumptions"
    ws4["A1"].font = TITLE_FONT
    assumptions = [
        ("Projection years", model.projection_years),
        ("Risk-free rate", model.risk_free_rate),
        ("Market risk premium", model.market_risk_premium),
        ("Beta", d.get("beta")),
        ("Terminal growth rate", model.terminal_growth_rate),
        ("Tax rate", model.tax_rate),
    ]
    ws4.cell(row=3, column=1, value="Assumption").font = BOLD
    ws4.cell(row=3, column=2, value="Value").font = BOLD
    _style_header_row(ws4, 3, 2)
    for i, (label, val) in enumerate(assumptions, start=1):
        ws4.cell(row=3 + i, column=1, value=label).border = BORDER
        c = ws4.cell(row=3 + i, column=2, value=val)
        c.border = BORDER
        if label in {"Risk-free rate", "Market risk premium", "Terminal growth rate", "Tax rate"}:
            c.number_format = "0.00%"
    _autofit(ws4, [26, 16])

    wb.save(filename)
    print(f"Excel workbook saved to: {filename}")


# --------------------------------------------------------------------------- #
# CLI entry point
# --------------------------------------------------------------------------- #

def main():
    parser = argparse.ArgumentParser(description="Automated DCF valuation tool with live data + Excel export.")
    parser.add_argument("ticker", help="Stock ticker symbol, e.g. AAPL, MSFT, GOOGL")
    parser.add_argument("--years", type=int, default=5, help="Number of years to project FCF (default 5)")
    parser.add_argument("--growth", type=float, default=None,
                         help="Manual FCF growth rate override, e.g. 0.08 for 8%% (default: auto-estimated from history)")
    parser.add_argument("--terminal-growth", type=float, default=0.025,
                         help="Terminal (perpetuity) growth rate (default 0.025)")
    parser.add_argument("--risk-free-rate", type=float, default=0.043,
                         help="Risk-free rate for CAPM, e.g. 10Y treasury yield (default 0.043)")
    parser.add_argument("--market-risk-premium", type=float, default=0.055,
                         help="Equity market risk premium (default 0.055)")
    parser.add_argument("--tax-rate", type=float, default=None,
                         help="Manual effective tax rate override (default: auto-estimated)")
    parser.add_argument("--output", default=None, help="Output .xlsx filename")
    args = parser.parse_args()

    model = DCFModel(
        ticker=args.ticker,
        projection_years=args.years,
        fcf_growth_rate=args.growth,
        terminal_growth_rate=args.terminal_growth,
        risk_free_rate=args.risk_free_rate,
        market_risk_premium=args.market_risk_premium,
        tax_rate=args.tax_rate,
    )

    try:
        model.run()
    except Exception as e:
        print(f"Error: {e}", file=sys.stderr)
        sys.exit(1)

    output_file = args.output or f"{model.ticker_symbol}_DCF_Valuation.xlsx"
    export_to_excel(model, output_file)


if __name__ == "__main__":
    main()
